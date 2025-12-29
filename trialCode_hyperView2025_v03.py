
# export_modal_gifs_no_tcl.py
# HyperView 2025: Export modal GIFs from Abaqus ODB
# NO Tcl/HWI; assemblies chosen via Excel list and applied with HWC show/hide.

import os
import time
import hw
import hw.hv as hv

# ---- USER SETTINGS (edit or set as environment vars) ----
ODB_PATH   = os.environ.get("HV_ODB",   r"K:\cae\powertrain\engedd\PT411\3.0_Programmes\Batteries\EMA_BEV\26MY\7-IPB\3-PackLevel\000150-PSD\231705-FullPackPSD\b1000\IPB4_R5p1_v0\X-DIR\Abaqus_res_001\EI_T9U3-10B759-AA15_EMA_LR_26MY_IPB4_R5P1_X_DIR_ov2023.odb")
MODE_LIST  = os.environ.get("HV_MODELIST", r"K:\cae\user\ykumar4\exchange\EI_T9U3-10B759-AA15_XCompModes.xlsx")   # Excel (.xlsx) with column "Mode"
ASM_XLSX      = os.environ.get("HV_ASM_LIST", r"K:\cae\user\ykumar4\exchange\assemblies.xlsx")  # column 'Assembly'
OUT_DIR       = os.environ.get("HV_OUTDIR", r"K:\cae\user\ykumar4\exchange\export_gifs_v03")

GIF_W, GIF_H  = 1920, 1080
PHASE_INC_DEG = 8                # 8° → 45 steps per modal cycle
VIEW_CMD      = "view orientation iso"  # vantage point

# Deformation scaling (optional; pick one line if you need)
SCALE_FACTOR  = 5      # exaggerate 5× actual displacement

# ---- helpers: read Excel sheets ----
def read_column_excel(path, header_name):
    if not os.path.isfile(path):
        raise FileNotFoundError(f"Excel file not found: {path}")
    try:
        import openpyxl
    except ImportError:
        raise ImportError("openpyxl not available. Install it or convert to CSV.")
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    header = [c.value for c in ws[1]]
    col = None
    for i, name in enumerate(header, start=1):
        if isinstance(name, str) and name.strip().lower() == header_name.lower():
            col = i; break
    if col is None:
        raise ValueError(f"Excel must have a header column named '{header_name}'.")
    vals = []
    for r in ws.iter_rows(min_row=2):
        v = r[col-1].value
        if v is None: continue
        vals.append(str(v).strip())
    return [x for x in vals if x]

def apply_assembly_visibility_from_excel():
    """NO Tcl: use HWC to hide all assemblies, then show the list from Excel."""
    asm_labels = read_column_excel(ASM_XLSX, "Assembly")
    if not asm_labels:
        raise ValueError(f"No assembly labels found in {ASM_XLSX} (column 'Assembly').")
    # Clear visibility: show only selected assemblies via labels/wildcards.  
    hw.evalHWC("hide assembly all")
    for label in asm_labels:
        # Wildcards work (e.g., '*', '?') and comma-separated label lists are supported. 
        hw.evalHWC(f'show assembly "{label}"')

def clear_and_disable_contours():
    """Remove any contour plots and ensure contour display is off."""
    # Clear the current scalar plot (restores original state). 
    hw.evalHWC('result scalar clear')
    # Optional (strong off): turn off contour display option (no Tcl variant—not strictly needed if 'clear' is used).
    # If you prefer stronger guarantee without Tcl, just keep 'clear'; legends can be hidden too:
    hw.evalHWC('hide legends')  # hides legends if any were visible. 

def main():
    os.makedirs(OUT_DIR, exist_ok=True)

    # New session and HyperView window
    ses = hw.Session(); ses.new()
    win = ses.get(hw.Window); win.type = "animation"  # HyperView client 

    # Load ODB as both model & result (HyperView can read geometry & results from ODB). 
    win.addModelAndResult(ODB_PATH, result=ODB_PATH)

    # View
    hw.evalHWC(VIEW_CMD)

    # Make sure contours are gone
    clear_and_disable_contours()

    # --- Display chosen assemblies (NO Tcl): read labels from Excel and show via HWC ---
    apply_assembly_visibility_from_excel()

    # --- Prepare modal animation & GIF capture ---
    res = ses.get(hv.Result)
    sim_ids = res.getSimulationIds()
    if not sim_ids:
        raise RuntimeError("No simulations found in ODB. Check that the ODB contains modal results.")

    modes = read_column_excel(MODE_LIST, "Mode")
    # Keep only integers; sort and dedupe
    try:
        modes = sorted(set(int(m) for m in modes))
    except Exception:
        raise ValueError("Mode list contains non-integer values.")

    # Capture tool
    cvt = hw.CaptureVideoTool()
    cvt.type = "gif"
    cvt.dimension = "pixels"
    cvt.width = GIF_W
    cvt.height = GIF_H

    # Modal animation behavior
    hw.evalHWC('hwd page current animationmode=modal')  # switch page to modal animation mode 

    anim = hw.AnimationTool()

    # Loop through modes
    for mode_idx in modes:
        if mode_idx < 1 or mode_idx > len(sim_ids):
            print(f"[WARN] Mode {mode_idx} out of range (1..{len(sim_ids)}). Skipping.")
            continue

        sim_id = sim_ids[mode_idx - 1]

        # Select current simulation (the chosen eigenmode) and configure oscillation.
        hw.evalHWC(f"result simulation {sim_id}")           # set mode by simulation ID 
        hw.evalHWC(f"animate modal startframe {sim_id}")    # oscillate this mode         
        hw.evalHWC(f"animate modal endframe {sim_id}")
        hw.evalHWC(f"animate modal increment {PHASE_INC_DEG}")
        hw.evalHWC(f'scale deformed resulttype="Displacement" scale=scalefactor value={SCALE_FACTOR}')    

        win.draw()
        out_file = os.path.join(OUT_DIR, f"modal_mode_{mode_idx:03d}.gif")
        cvt.file = out_file
        print(f"[INFO] Capturing GIF for mode {mode_idx} → {out_file}")
        cvt.capture()

        # small live feedback
        anim.start(num_cycles=0)
        time.sleep(2)
        anim.stop()

    print(f"[DONE] GIFs exported → {OUT_DIR}")

if __name__ == "__main__":
    try:
        main()
    except Exception as e:
        print("[ERROR]", e)
        raise
